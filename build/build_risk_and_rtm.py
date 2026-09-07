"""Risk Register (Document 09) and Requirements Traceability Matrix (Document 15)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"; NAVY = "1E2761"; ICE = "EEF2FA"
thin = Side(style="thin", color="C9CFDD"); BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)
HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF"); BODY = Font(name=FONT, size=10); BLUE = Font(name=FONT, size=10, color="0000FF")
TITLE = Font(name=FONT, size=14, bold=True, color=NAVY); SUB = Font(name=FONT, size=10, italic=True, color="5A6270")


def sheet(wb, name, title, subtitle, widths):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = TITLE
    ws["A2"] = subtitle; ws["A2"].font = SUB
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws


def header(ws, row, labels):
    for i, l in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=l)
        c.font = HDR; c.fill = PatternFill("solid", fgColor=NAVY); c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BORDER
    ws.row_dimensions[row].height = 32


def cell(ws, r, c, v, font=BODY, fmt=None, fill=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = font; x.border = BORDER; x.alignment = Alignment(wrap_text=True, vertical="top")
    if fmt: x.number_format = fmt
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    return x


# ================================================================== RISK REGISTER
wb = Workbook(); wb.remove(wb.active)
ws = sheet(wb, "Risk_Register", "Risk Register", "Score = Likelihood x Impact (1-5 each). Rating: 15+ High, 8-14 Medium, under 8 Low. Blue cells are inputs; update at each steering call.",
           [7, 16, 44, 11, 9, 8, 10, 52, 22, 34, 12])
header(ws, 4, ["ID", "Category", "Risk", "Likelihood (1-5)", "Impact (1-5)", "Score", "Rating", "Mitigation", "Owner", "Trigger / early indicator", "Status"])
risks = [
    ("Data and compliance", "Client's legal team requires that no document text leaves the country, making Pattern A unusable for the pilot", 3, 4, "Settle residency in discovery (week 1); gateway supports in-region routing; sovereign pattern designed and priced; keep a small in-region model available for restricted classes", "Engagement lead", "Legal review outcome at discovery", "Open"),
    ("Data and compliance", "Personal data of investors and contacts stored or contacted without a lawful basis under the Malaysian PDPA or the Philippine Data Privacy Act", 3, 5, "Consent and lawful-basis fields per contact; outreach flag when consent missing; data protection checkpoint before real data; legal confirms basis per jurisdiction; DPO appointed as the amended PDPA requires", "Client legal + PM", "Missing consent rate in import preview", "Open"),
    ("Data and compliance", "Frontier model availability or terms in the Southeast Asian regions change during the project (cross-region inference, retention terms)", 3, 3, "Model gateway abstracts providers; re-verify availability at discovery and before go-live; open-weight fallback evaluated in Phase 3", "AI engineer", "Provider announcements; failed region checks", "Open"),
    ("Accuracy", "Claims extraction precision below target on messy scanned or non-English documents", 4, 4, "Sample corpus in week 1 includes worst cases; OCR with layout analysis; confidence thresholds route low-confidence extractions to human review; tune prompts and tolerances in pilot", "AI engineer", "Labelled-sample precision below 90% at week 6", "Open"),
    ("Accuracy", "False-positive contradiction flags erode trust (superseded drafts, rounding, currency)", 4, 3, "Version detection and supersedes links; numeric tolerance and currency normalisation; adjudication pass with rationale; dismiss-with-reason feeds tuning; both quotes always shown", "AI engineer", "False-positive rate above 15% in pilot", "Open"),
    ("Accuracy", "Answers hallucinate or cite the wrong page", 2, 5, "Grounding check removes unsupported sentences; refusal when nothing survives; citation accuracy in golden set; regression gate in CI", "AI engineer", "Citation accuracy below 95%", "Open"),
    ("Scope", "Decision types and evidence checklists not defined by the client, so gap tracking cannot be demonstrated", 3, 4, "Two discovery sessions dedicated to decision types; start with five; template library of common decisions as a fallback", "Product owner", "No catalogue by end of week 2", "Open"),
    ("Scope", "Scope creep: additional connectors, ERP integration, mobile app requested during pilot", 4, 3, "Change control with impact on time and cost; Phase 5 roadmap holds requests; sponsor approves changes", "PM", "Requests outside SOW at steering", "Open"),
    ("Delivery", "Late client inputs (sample documents, identity access, mailbox consent) delay milestones", 4, 3, "Dependency list with dates; weekly steering; timeline extends by the delay per SOW; IT contact named at kick-off", "PM + client IT", "Dependency overdue by 5 days", "Open"),
    ("Delivery", "Key engineer unavailable mid-project", 2, 4, "Documented code and architecture; two engineers share knowledge from week 5; supplier bench cover", "Engagement lead", "Absence notice", "Open"),
    ("Delivery", "Underestimated effort for multilingual search and interface once the true language mix of the corpus is known", 3, 3, "Language mix confirmed in week 1 from the sample corpus; multilingual retrieval in scope from Phase 1, not retrofitted; per-language evaluation and reporting", "Full-stack engineer", "Any language's golden-set scores lag English by more than 10 points", "Open"),
    ("Security", "Prompt injection via a malicious or crafted document changes assistant behaviour", 2, 4, "Retrieved text treated as data with delimiters; injection-pattern flags at ingestion; no tool actions from the assistant; penetration test includes injection cases", "Security specialist", "Injection test failures", "Open"),
    ("Security", "Access control mis-configuration exposes restricted documents to unauthorised users", 2, 5, "Row-level security in database; ACL mirrored from source; permission tests in CI; access review at go-live", "Security specialist", "Failed permission test; audit anomaly", "Open"),
    ("Security", "Staff bypass the approval gate by sending from personal mailboxes", 3, 2, "Policy plus reporting of approved-route sends; connected mailbox captures replies; management review of outreach volumes", "Client compliance", "Outreach volumes drop while replies rise", "Open"),
    ("Cost", "AI usage costs exceed pilot budget due to heavy re-indexing or long documents", 2, 2, "Gateway budgets and alerts; batch and caching on ingestion; Sonnet for routine questions; monthly usage report", "AI engineer", "Spend above 120% of forecast", "Open"),
    ("Cost", "Sovereign GPU hosting cost higher than modelled or capacity unavailable in region", 3, 3, "Reserved pricing; right-size model (smaller open-weight model for routine tasks); decide after pilot with measured usage", "Engagement lead", "Quotes above model by more than 25%", "Open"),
    ("Adoption", "Relationship managers do not log interactions, so reporting and reminders are unreliable", 4, 4, "Pilot champions; mailbox and calendar capture reduce manual logging; simple mobile-friendly forms; management uses the dashboards visibly", "Product owner", "Interactions logged per user per week below target", "Open"),
    ("Adoption", "Executives lose trust after an early wrong answer", 3, 4, "Set expectations: cited, checkable answers; show refusals as a feature; weekly accuracy reporting; feedback button on every answer", "Engagement lead", "Usage drop after an incident", "Open"),
    ("Vendor", "Third-party service (embedding, OCR, hosting) price or terms change", 2, 2, "All components swappable; alternatives listed in Document 06; contracts in the client's name", "PM", "Vendor notice", "Open"),
    ("Commercial", "Client expects a single fixed price before discovery", 3, 3, "Present tiers and structure; offer fixed price for discovery and demo; workbook transparency", "Engagement lead", "Procurement asks for a total before week 0", "Open"),
    ("Operational", "Production incident without an on-call arrangement in the departmental tier", 2, 3, "Support tier chosen before go-live; runbooks; monitoring alerts to a named rota", "DevOps specialist", "Incident outside business hours", "Open"),
    ("Operational", "Backups not restorable when needed", 1, 5, "Quarterly restore tests; point-in-time recovery; cross-region copy at enterprise tier", "DevOps specialist", "Failed restore test", "Open"),
]
dv_l = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=False)
dv_s = DataValidation(type="list", formula1='"Open,Mitigating,Closed,Accepted"', allow_blank=False)
ws.add_data_validation(dv_l); ws.add_data_validation(dv_s)
for i, (cat, risk, l, im, mit, owner, trig, status) in enumerate(risks, start=5):
    cell(ws, i, 1, f"R{i - 4:02d}"); cell(ws, i, 2, cat); cell(ws, i, 3, risk)
    cell(ws, i, 4, l, BLUE); cell(ws, i, 5, im, BLUE)
    cell(ws, i, 6, f"=D{i}*E{i}"); cell(ws, i, 7, f'=IF(F{i}>=15,"High",IF(F{i}>=8,"Medium","Low"))')
    cell(ws, i, 8, mit); cell(ws, i, 9, owner); cell(ws, i, 10, trig); cell(ws, i, 11, status, BLUE)
    dv_l.add(f"D{i}"); dv_l.add(f"E{i}"); dv_s.add(f"K{i}")
last = 4 + len(risks)
ws.conditional_formatting.add(f"G5:G{last}", CellIsRule(operator="equal", formula=['"High"'], fill=PatternFill("solid", fgColor="F4C7C3")))
ws.conditional_formatting.add(f"G5:G{last}", CellIsRule(operator="equal", formula=['"Medium"'], fill=PatternFill("solid", fgColor="FCE8B2")))
ws.conditional_formatting.add(f"G5:G{last}", CellIsRule(operator="equal", formula=['"Low"'], fill=PatternFill("solid", fgColor="D9EAD3")))
ws.freeze_panes = "C5"
r = last + 2
cell(ws, r, 1, "Summary"); ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)
cell(ws, r + 1, 1, "High"); cell(ws, r + 1, 2, f'=COUNTIF(G5:G{last},"High")')
cell(ws, r + 2, 1, "Medium"); cell(ws, r + 2, 2, f'=COUNTIF(G5:G{last},"Medium")')
cell(ws, r + 3, 1, "Low"); cell(ws, r + 3, 2, f'=COUNTIF(G5:G{last},"Low")')
cell(ws, r + 4, 1, "Open"); cell(ws, r + 4, 2, f'=COUNTIF(K5:K{last},"Open")')
cell(ws, r + 5, 1, "Average score"); cell(ws, r + 5, 2, f"=AVERAGE(F5:F{last})", fmt="0.0")

ws2 = sheet(wb, "Scoring_Guide", "Scoring guide", "Use these definitions to keep scores consistent across reviewers.", [12, 40, 40])
header(ws2, 4, ["Score", "Likelihood", "Impact"])
guide = [(1, "Rare: not expected in this project", "Negligible: absorbed within the week"), (2, "Unlikely: could occur", "Minor: under one week delay or under 2% cost"), (3, "Possible: may well occur", "Moderate: one to three weeks delay or 2-8% cost"), (4, "Likely: expected at least once", "Major: milestone slip or 8-20% cost; trust affected"), (5, "Almost certain", "Severe: project viability, regulatory breach, or data exposure")]
for i, (s, l, im) in enumerate(guide, start=5):
    cell(ws2, i, 1, s); cell(ws2, i, 2, l); cell(ws2, i, 3, im)
wb.save("/home/user/CRM/proposal/09_Risk_Register.xlsx"); print("wrote risk register")

# ================================================================== TRACEABILITY MATRIX
wb = Workbook(); wb.remove(wb.active)
ws = sheet(wb, "Traceability", "Requirements Traceability Matrix", "Every requirement in the client's brief traced to design, component, phase, acceptance test and demo step. Status is an input.",
           [8, 34, 40, 26, 24, 12, 44, 16, 12])
header(ws, 4, ["Req ID", "Requirement (client brief)", "Design response (Document 03)", "Component", "Phase / milestone", "Doc 04 acceptance", "Acceptance test", "Demo step (Doc 12)", "Status"])
reqs = [
    ("R1", "Retrieve specific info from documents on request; controlled access; not open internet", "Hybrid retrieval + re-ranking over private corpus; permission filter applied before retrieval; closed-domain system instruction (s.9.2)", "Intelligence service, model gateway, PostgreSQL/pgvector", "Phase 1 / M1", "s.6 Retrieval", "Golden set: source in top results 90%+; permission tests: user without ACL never receives passage", "3.2 step 2", "Planned"),
    ("R2", "Flag contradictions: different figures, dates or terms for the same thing across documents", "Claims ledger with normalisation; cross-document comparison with tolerance; LLM adjudication for terms; findings with both quotes (s.3.3, s.9.1 steps 6-7)", "Intelligence service, claims and findings tables", "Phase 2 / M2", "s.6 Contradiction detection", "Seeded conflicts detected 90%+; false positives under 15%; superseded versions not flagged", "3.2 steps 2-3; seeded issues 1-3", "Planned"),
    ("R3", "Track missing information: gaps where a decision needs a document or data point not yet provided", "Decision checklists (configurable) matched against ledger and inventory; dangling-reference detection (s.9.1 step 8)", "Intelligence service, decisions and checklist tables", "Phase 2 / M2", "s.6 Gap tracking", "All catalogued checklist items evaluated correctly on test decisions; dangling references detected", "3.2 step 4; seeded issues 4-5", "Planned"),
    ("R4", "Answer management questions in plain language based only on the loaded document set", "Grounded generation with citations; independent grounding check; refusal when unsupported; multilingual (s.9.2 steps 5-6)", "Intelligence service, model gateway, web app", "Phase 1 / M1", "s.6 Answer faithfulness; Refusal correctness", "Faithfulness 95%+; refusal correctness 90%+; per-language reporting", "3.2 steps 2, 5, 6; seeded issues 6, 8", "Planned"),
    ("R5", "Closed-domain RAG with verification/audit layer, not a general chatbot", "Audit log of every query with retrieved passages and answer; hash-chained events; findings workbench (s.3.2, s.12)", "Platform audit log, web app", "Phase 1 / M1", "s.6 Security", "Audit chain verifies; every query has a logged trace", "3.2 step 11", "Planned"),
    ("C1", "Three categories: investors; partners/counterparties; customers/members; GCC-specific pipeline", "Category-typed relationships; configurable pipelines and stage sets per category; GCC template for market entry (s.2.2, s.8.3)", "Core API, CRM schema, web app", "Phase 1 / M1", "s.6 Reporting (reconciliation)", "All three categories and the GCC pipeline demonstrated; stages configurable without code", "3.2 step 7, 10", "Planned"),
    ("C2", "Contact and relationship records with stage/status tracking", "Organisation, Person, Relationship entities; stage history; status values incl. dormant (s.8.3)", "Entity registry, CRM schema", "Phase 1 / M1", "s.6 Reporting", "Record CRUD and stage transitions with history; import of 500-row sample", "3.2 step 7", "Planned"),
    ("C3", "Follow-up scheduling and reminders", "Follow-up tasks with due dates and owners; daily and hourly reminder job to staff; overdue and stale surfacing (s.9.4)", "Workers, Core API, web app", "Phase 1 / M1", "s.6 Reporting", "Reminder generated for due and overdue follow-ups; none sent to external contacts", "3.2 step 7", "Planned"),
    ("C4", "Any outreach must go through human approval before sending; no fully automated outbound", "Outreach Request workflow; four-eyes rule enforced in database; no scheduled-send path; send under approver's authority via connected mailbox; versioned and logged (s.9.3, s.12.3)", "Core API, approval policies, mailbox connector, audit log", "Phase 2 / M2", "s.6 Approval gate", "Automated tests: send without approval fails; drafter cannot approve; audit query shows approver on every send", "3.2 step 8", "Planned"),
    ("C5", "Basic reporting: who has been contacted, stage, overdue", "Dashboards and exports: activity by period, pipeline by stage and category, overdue, dormant, approval turnaround (s.10.1)", "Core API reports, web app", "Phase 2 / M2", "s.6 Reporting", "Dashboard figures reconcile to database counts", "3.2 step 9", "Planned"),
    ("X1", "Proposed build approach and tools for each system", "Documents 03 and 06", "n/a", "Proposal", "n/a", "Client review", "First meeting slides 6-12", "Delivered"),
    ("X2", "Structure that accepts real data later without rebuild", "Schema-first with custom fields; configuration-driven rules; import templates; connectors; removable demo data (s.5; Document 07)", "Platform, Core API, web app admin", "Phase 1-2 / M2", "s.6 Reporting (import)", "Custom field and stage added without code; import preview; demo-data removal leaves configuration intact", "3.2 step 10", "Planned"),
    ("X3", "Timeline to working structure and demo", "Document 04: working structure week 6; integrated demo week 8-9", "n/a", "Proposal", "n/a", "Client review", "Slide 14", "Delivered"),
    ("X4", "Cost or pricing model", "Documents 05 and 14: run-cost tiers, build effort, three pricing options, 3-year TCO", "n/a", "Proposal", "n/a", "Client review", "Slides 15-16", "Delivered"),
    ("N1", "Documents, questions, answers and interface in the languages the corpus actually uses; English primary (mix to be confirmed; Document 13 gap 1)", "OCR configured per language; multilingual embeddings; localisable interface with RTL available if later needed; per-language evaluation", "Intelligence service, web app", "Phase 1-3", "s.6 Answer faithfulness", "Every in-scope language within 10 points of English on the golden set", "3.2 step 5", "Planned"),
    ("N2", "Data residency by policy (Document 13 gap 2)", "Deployment patterns A/B/C; gateway routing by classification (s.3.5)", "Model gateway, classification step", "Phase 0 decision; Phase 1 build", "s.6 Security", "Restricted-class document never routed to external provider (gateway log test)", "Slide 11", "Planned"),
    ("N3", "Document version handling (Document 13 gap 5)", "Checksum and naming detection; supersedes links; comparison excludes superseded (s.9.1 step 1)", "Intelligence service", "Phase 2 / M2", "s.6 Contradiction detection", "Seeded v1/v2 pair produces no contradiction flag", "Seeded issue 3", "Planned"),
]
dv = DataValidation(type="list", formula1='"Planned,In progress,Delivered,Accepted,Deferred"', allow_blank=False); ws.add_data_validation(dv)
for i, row in enumerate(reqs, start=5):
    for j, v in enumerate(row, start=1):
        cell(ws, i, j, v, BLUE if j == 9 else BODY)
    dv.add(f"I{i}")
ws.freeze_panes = "C5"
last = 4 + len(reqs)
r = last + 2
cell(ws, r, 1, "Coverage summary"); ws.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)
for k, st in enumerate(["Planned", "In progress", "Delivered", "Accepted", "Deferred"]):
    cell(ws, r + 1 + k, 1, st); cell(ws, r + 1 + k, 2, f'=COUNTIF(I5:I{last},"{st}")')
cell(ws, r + 6, 1, "Total"); cell(ws, r + 6, 2, f"=COUNTA(A5:A{last})")
wb.save("/home/user/CRM/proposal/15_Requirements_Traceability_Matrix.xlsx"); print("wrote RTM")
